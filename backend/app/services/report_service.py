from datetime import date
from sqlalchemy.orm import Session
from sqlalchemy import func

from backend.app.models.user import User
from backend.app.models.contract import Contract
from backend.app.models.obligation import Obligation
from backend.app.models.renewal import Renewal
from backend.app.models.compliance import Compliance
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from io import BytesIO
from reportlab.lib.pagesizes import A4
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle
)
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet

# ============================================================
# CONTRACT ANALYTICS
# ============================================================

def get_contract_summary(db: Session):

    total = db.query(func.count(Contract.id)).scalar() or 0

    active = db.query(func.count(Contract.id)).filter(
        Contract.status == "Active"
    ).scalar() or 0

    draft = db.query(func.count(Contract.id)).filter(
        Contract.status == "Draft"
    ).scalar() or 0

    under_review = db.query(func.count(Contract.id)).filter(
        Contract.status == "Under Review"
    ).scalar() or 0

    approved = db.query(func.count(Contract.id)).filter(
        Contract.status == "Approved"
    ).scalar() or 0

    expired = db.query(func.count(Contract.id)).filter(
        Contract.status == "Expired"
    ).scalar() or 0

    terminated = db.query(func.count(Contract.id)).filter(
        Contract.status == "Terminated"
    ).scalar() or 0
    category_data = db.query(
    Contract.category,
    func.count(Contract.id)
    ).group_by(
    Contract.category
    ).all()


    categories = {}

    for category, count in category_data:
        categories[category or "Uncategorized"] = count
    return {
        "total": total,
        "active": active,
        "draft": draft,
        "under_review": under_review,
        "approved": approved,
        "expired": expired,
        "terminated": terminated,
        "categories": categories
}
    
    # ============================================================
# OBLIGATION ANALYTICS
# ============================================================

from backend.app.models.obligation import Obligation


def get_obligation_summary(db: Session):

    total = db.query(
        func.count(Obligation.id)
    ).scalar() or 0

    pending = db.query(
        func.count(Obligation.id)
    ).filter(
        Obligation.status == "Pending"
    ).scalar() or 0

    in_progress = db.query(
        func.count(Obligation.id)
    ).filter(
        Obligation.status == "In Progress"
    ).scalar() or 0

    completed = db.query(
        func.count(Obligation.id)
    ).filter(
        Obligation.status == "Completed"
    ).scalar() or 0

    delayed = db.query(
        func.count(Obligation.id)
    ).filter(
        Obligation.status == "Delayed"
    ).scalar() or 0

    overdue = db.query(
        func.count(Obligation.id)
    ).filter(
        Obligation.status == "Overdue"
    ).scalar() or 0

    return {
        "total": total,
        "pending": pending,
        "in_progress": in_progress,
        "completed": completed,
        "delayed": delayed,
        "overdue": overdue
    }
    # ============================================================
# RENEWAL ANALYTICS
# ============================================================

def get_renewal_summary(db: Session):

    total_upcoming = db.query(
        func.count(Renewal.id)
    ).filter(
        Renewal.status == "Upcoming"
    ).scalar() or 0

    in_progress = db.query(
        func.count(Renewal.id)
    ).filter(
        Renewal.status == "In Progress"
    ).scalar() or 0

    renewed = db.query(
        func.count(Renewal.id)
    ).filter(
        Renewal.status == "Renewed"
    ).scalar() or 0

    expired = db.query(
        func.count(Renewal.id)
    ).filter(
        Renewal.status == "Expired"
    ).scalar() or 0

    cancelled = db.query(
        func.count(Renewal.id)
    ).filter(
        Renewal.status == "Cancelled"
    ).scalar() or 0

    return {
        "upcoming": total_upcoming,
        "in_progress": in_progress,
        "renewed": renewed,
        "expired": expired,
        "cancelled": cancelled
    }
    # ============================================================
# COMPLIANCE ANALYTICS
# ============================================================

def get_compliance_summary(db: Session):

    total_contracts = db.query(
        func.count(Contract.id)
    ).scalar() or 0

    compliant = 0
    pending = 0
    delayed = 0
    non_compliant = 0
    high_risk = 0

    total_score = 0
    evaluated_count = 0

    contracts = db.query(Contract).all()

    for contract in contracts:

        latest_compliance = db.query(Compliance).filter(
            Compliance.contract_id == contract.id
        ).order_by(
            Compliance.evaluated_at.desc()
        ).first()

        if latest_compliance:

            evaluated_count += 1

            if latest_compliance.compliance_score:
                total_score += latest_compliance.compliance_score

            if latest_compliance.status == "Compliant":
                compliant += 1

            elif latest_compliance.status == "Pending":
                pending += 1

            elif latest_compliance.status == "Delayed":
                delayed += 1

            elif latest_compliance.status == "Non-Compliant":
                non_compliant += 1

            if latest_compliance.risk_level == "High":
                high_risk += 1

    average_compliance_score = 0

    if evaluated_count > 0:

        average_compliance_score = round(
            total_score / evaluated_count,
            2
        )

    return {
        "total_contracts": total_contracts,
        "compliant": compliant,
        "pending": pending,
        "delayed": delayed,
        "non_compliant": non_compliant,
        "high_risk": high_risk,
        "average_compliance_score": average_compliance_score
    }
    # ============================================================
# RISK ANALYSIS
# ============================================================

def get_risk_report(db: Session):

    contracts = db.query(Contract).all()

    results = []

    for contract in contracts:

        latest_compliance = db.query(Compliance).filter(
            Compliance.contract_id == contract.id
        ).order_by(
            Compliance.evaluated_at.desc()
        ).first()

        if latest_compliance:

            if latest_compliance.risk_level in ["Medium", "High"]:

                overdue_obligations = db.query(
                    func.count(Obligation.id)
                ).filter(
                    Obligation.contract_id == contract.id,
                    Obligation.status == "Overdue"
                ).scalar() or 0

                results.append({
                    "contract_id": contract.id,
                    "contract_number": contract.contract_number,
                    "risk_level": latest_compliance.risk_level,
                    "overdue_obligations": overdue_obligations,
                    "compliance_score": latest_compliance.compliance_score or 0
                })

    return results
# ============================================================
# CONTRACT PDF REPORT
# ============================================================

def generate_contract_pdf(db: Session):

    contracts = db.query(Contract).all()

    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4
    )

    elements = []

    styles = getSampleStyleSheet()

    title = Paragraph(
        "ContractIQ - Contract Report",
        styles["Title"]
    )

    elements.append(title)
    elements.append(Spacer(1, 20))

    data = [
        [
            "Contract Number",
            "Title",
            "Category",
            "Status",
            "Start Date",
            "End Date",
            "Assigned User"
        ]
    ]

    for contract in contracts:

        assigned_user = ""

        if contract.assigned_user:
            assigned_user = contract.assigned_user.email or ""

        data.append([
            str(contract.contract_number or ""),
            str(contract.title or ""),
            str(contract.category or ""),
            str(contract.status or ""),
            str(contract.start_date or ""),
            str(contract.end_date or ""),
            str(assigned_user)
        ])

    table = Table(
        data,
        repeatRows=1
    )

    table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 1, colors.black),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
        ])
    )

    elements.append(table)

    doc.build(elements)

    buffer.seek(0)

    return buffer
# ============================================================
# CONTRACT EXCEL REPORT
# ============================================================

def generate_contract_excel(db: Session):

    contracts = db.query(Contract).all()

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Contracts"

    headers = [
        "Contract Number",
        "Title",
        "Category",
        "Status",
        "Start Date",
        "End Date",
        "Assigned User"
    ]

    worksheet.append(headers)


    # Header formatting

    for cell in worksheet[1]:

        cell.font = Font(bold=True)

        cell.alignment = Alignment(
            horizontal="center"
        )


    # Contract data

    for contract in contracts:

        assigned_user = ""

        if contract.assigned_user:

            assigned_user = (
                contract.assigned_user.email or ""
            )

        worksheet.append([
            contract.contract_number or "",
            contract.title or "",
            contract.category or "",
            contract.status or "",
            str(contract.start_date or ""),
            str(contract.end_date or ""),
            assigned_user
        ])


    # Column widths

    worksheet.column_dimensions["A"].width = 20

    worksheet.column_dimensions["B"].width = 30

    worksheet.column_dimensions["C"].width = 20

    worksheet.column_dimensions["D"].width = 18

    worksheet.column_dimensions["E"].width = 15

    worksheet.column_dimensions["F"].width = 15

    worksheet.column_dimensions["G"].width = 30


    buffer = BytesIO()

    workbook.save(buffer)

    buffer.seek(0)

    return buffer
# ============================================================
# OBLIGATION PDF REPORT
# ============================================================

def generate_obligation_pdf(db: Session):

    obligations = db.query(Obligation).all()

    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4
    )

    elements = []

    styles = getSampleStyleSheet()

    title = Paragraph(
        "ContractIQ - Obligation Report",
        styles["Title"]
    )

    elements.append(title)
    elements.append(Spacer(1, 20))

    data = [
        [
            "ID",
            "Contract ID",
            "Title",
            "Due Date",
            "Status"
        ]
    ]

    for obligation in obligations:

        data.append([
            str(obligation.id or ""),
            str(obligation.contract_id or ""),
            str(obligation.title or ""),
            str(obligation.due_date or ""),
            str(obligation.status or "")
        ])

    table = Table(
        data,
        repeatRows=1
    )

    table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 1, colors.black),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
        ])
    )

    elements.append(table)

    doc.build(elements)

    buffer.seek(0)

    return buffer
# ============================================================
# OBLIGATION EXCEL REPORT
# ============================================================

def generate_obligation_excel(db: Session):

    obligations = db.query(Obligation).all()

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Obligations"

    headers = [

        "Contract",

        "Obligation Title",

        "Obligation Type",

        "Assigned User",

        "Due Date",

        "Status",

        "Completion Date"

    ]

    worksheet.append(headers)


    # Header Formatting

    for cell in worksheet[1]:

        cell.font = Font(bold=True)

        cell.alignment = Alignment(
            horizontal="center"
        )


    # Obligation Data

    for obligation in obligations:

        contract_number = ""

        if obligation.contract:

            contract_number = (
                obligation.contract.contract_number or ""
            )


        assigned_user = ""

        if obligation.user:

            assigned_user = (
                obligation.user.email or ""
            )


        worksheet.append([

            contract_number,

            obligation.title or "",

            obligation.obligation_type or "",

            assigned_user,

            str(obligation.due_date or ""),

            obligation.status or "",

            str(obligation.completion_date or "")

        ])


    # Column Widths

    worksheet.column_dimensions["A"].width = 20

    worksheet.column_dimensions["B"].width = 30

    worksheet.column_dimensions["C"].width = 20

    worksheet.column_dimensions["D"].width = 30

    worksheet.column_dimensions["E"].width = 18

    worksheet.column_dimensions["F"].width = 18

    worksheet.column_dimensions["G"].width = 20


    buffer = BytesIO()

    workbook.save(buffer)

    buffer.seek(0)

    return buffer
# ============================================================
# RENEWAL PDF REPORT
# ============================================================

def generate_renewal_pdf(db: Session):

    renewals = db.query(Renewal).all()

    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4
    )

    elements = []

    styles = getSampleStyleSheet()

    title = Paragraph(
        "ContractIQ - Renewal Report",
        styles["Title"]
    )

    elements.append(title)
    elements.append(Spacer(1, 20))

    data = [
        [
            "Contract",
            "Previous Expiry",
            "Renewal Date",
            "New Expiry",
            "Status",
            "Assigned User"
        ]
    ]

    for renewal in renewals:

        contract_number = ""

        if renewal.contract:
            contract_number = (
                renewal.contract.contract_number or ""
            )

        assigned_user = ""

        if renewal.assigned_user:
            assigned_user = (
                renewal.assigned_user.email or ""
            )

        data.append([
            str(contract_number),
            str(renewal.previous_expiry_date or ""),
            str(renewal.renewal_date or ""),
            str(renewal.new_expiry_date or ""),
            str(renewal.status or ""),
            str(assigned_user)
        ])

    table = Table(
        data,
        repeatRows=1
    )

    table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 1, colors.black),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
        ])
    )

    elements.append(table)

    doc.build(elements)

    buffer.seek(0)

    return buffer
# ============================================================
# RENEWAL EXCEL REPORT
# ============================================================

def generate_renewal_excel(db: Session):

    renewals = db.query(Renewal).all()

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Renewals"

    headers = [
        "Contract",
        "Previous Expiry Date",
        "Renewal Date",
        "New Expiry Date",
        "Renewal Status",
        "Assigned User"
    ]

    worksheet.append(headers)

    # Header formatting
    for cell in worksheet[1]:

        cell.font = Font(bold=True)

        cell.alignment = Alignment(
            horizontal="center"
        )

    # Renewal data
    for renewal in renewals:

        contract_number = ""

        if renewal.contract:
            contract_number = (
                renewal.contract.contract_number or ""
            )

        assigned_user = ""

        if renewal.assigned_user:
            assigned_user = (
                renewal.assigned_user.email or ""
            )

        worksheet.append([
            contract_number,
            str(renewal.previous_expiry_date or ""),
            str(renewal.renewal_date or ""),
            str(renewal.new_expiry_date or ""),
            renewal.status or "",
            assigned_user
        ])

    # Column widths
    worksheet.column_dimensions["A"].width = 20
    worksheet.column_dimensions["B"].width = 22
    worksheet.column_dimensions["C"].width = 20
    worksheet.column_dimensions["D"].width = 22
    worksheet.column_dimensions["E"].width = 20
    worksheet.column_dimensions["F"].width = 30

    buffer = BytesIO()

    workbook.save(buffer)

    buffer.seek(0)

    return buffer
# ============================================================
# COMPLIANCE PDF REPORT
# ============================================================

def generate_compliance_pdf(db: Session):

    compliances = db.query(Compliance).all()

    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4
    )

    elements = []

    styles = getSampleStyleSheet()

    title = Paragraph(
        "ContractIQ - Compliance Report",
        styles["Title"]
    )

    elements.append(title)
    elements.append(Spacer(1, 20))

    data = [
        [
            "Contract",
            "Status",
            "Score",
            "Overdue",
            "Risk Level",
            "Evaluation Date"
        ]
    ]

    for compliance in compliances:

        contract_number = ""

        if compliance.contract:
            contract_number = (
                compliance.contract.contract_number or ""
            )

        overdue_obligations = 0

        if compliance.contract_id:

            overdue_obligations = db.query(
                func.count(Obligation.id)
            ).filter(
                Obligation.contract_id == compliance.contract_id,
                Obligation.status == "Overdue"
            ).scalar() or 0

        data.append([
            str(contract_number),
            str(compliance.status or ""),
            str(compliance.compliance_score or 0),
            str(overdue_obligations),
            str(compliance.risk_level or ""),
            str(compliance.evaluated_at or "")
        ])

    table = Table(
        data,
        repeatRows=1
    )

    table.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 1, colors.black),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
        ])
    )

    elements.append(table)

    doc.build(elements)

    buffer.seek(0)

    return buffer
# ============================================================
# COMPLIANCE EXCEL REPORT
# ============================================================

def generate_compliance_excel(db: Session):

    compliances = db.query(Compliance).all()

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Compliance"

    headers = [
        "Contract",
        "Compliance Status",
        "Compliance Score",
        "Overdue Obligations",
        "Risk Level",
        "Evaluation Date"
    ]

    worksheet.append(headers)

    # Header formatting
    for cell in worksheet[1]:

        cell.font = Font(bold=True)

        cell.alignment = Alignment(
            horizontal="center"
        )

    # Compliance data
    for compliance in compliances:

        contract_number = ""

        if compliance.contract:

            contract_number = (
                compliance.contract.contract_number or ""
            )

        overdue_obligations = 0

        if compliance.contract_id:

            overdue_obligations = db.query(
                func.count(Obligation.id)
            ).filter(
                Obligation.contract_id == compliance.contract_id,
                Obligation.status == "Overdue"
            ).scalar() or 0

        worksheet.append([
            contract_number,
            compliance.status or "",
            compliance.compliance_score or 0,
            overdue_obligations,
            compliance.risk_level or "",
            str(compliance.evaluated_at or "")
        ])

    # Column widths
    worksheet.column_dimensions["A"].width = 20
    worksheet.column_dimensions["B"].width = 22
    worksheet.column_dimensions["C"].width = 20
    worksheet.column_dimensions["D"].width = 22
    worksheet.column_dimensions["E"].width = 18
    worksheet.column_dimensions["F"].width = 25

    buffer = BytesIO()

    workbook.save(buffer)

    buffer.seek(0)

    return buffer
# ============================================================
# DEPARTMENT PERFORMANCE
# ============================================================

def get_department_performance(db):

    from backend.app.models.user import User
    from backend.app.models.contract import Contract
    from backend.app.models.obligation import Obligation

    departments = db.query(User.department).distinct().all()

    results = []

    for department_row in departments:

        department = department_row[0]

        # Skip users without department
        if not department:
            continue

        users = db.query(User).filter(
            User.department == department
        ).all()

        user_ids = [user.id for user in users]

        contracts_count = db.query(Contract).filter(
            Contract.owner_id.in_(user_ids)
        ).count()

        obligations_count = db.query(Obligation).filter(
            Obligation.assigned_to.in_(user_ids)
        ).count()

        overdue_count = db.query(Obligation).filter(
            Obligation.assigned_to.in_(user_ids),
            Obligation.status == "Overdue"
        ).count()

        results.append({
            "department": department,
            "contracts": contracts_count,
            "obligations": obligations_count,
            "overdue": overdue_count
        })

    return results
# ============================================================
# UPCOMING RENEWALS
# ============================================================

def get_upcoming_renewals(db: Session):

    renewals = db.query(Renewal).filter(
        Renewal.status == "Upcoming"
    ).all()

    results = []

    for renewal in renewals:

        contract = db.query(Contract).filter(
            Contract.id == renewal.contract_id
        ).first()

        results.append({
            "renewal_id": renewal.id,
            "contract_id": renewal.contract_id,
            "contract_number": (
                contract.contract_number if contract else None
            ),
            "contract_title": (
                contract.title if contract else None
            ),
            "renewal_date": renewal.renewal_date,
            "status": renewal.status
        })

    return results
# ============================================================
# OVERDUE OBLIGATIONS
# ============================================================

def get_overdue_obligations(db: Session):

    obligations = db.query(Obligation).filter(
        Obligation.status == "Overdue"
    ).all()

    results = []

    for obligation in obligations:

        results.append({
            "id": obligation.id,
            "contract_id": obligation.contract_id,
            "title": obligation.title,
            "due_date": obligation.due_date,
            "status": obligation.status
        })

    return results