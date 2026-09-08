from datetime import date, timedelta
from io import BytesIO

from fastapi import HTTPException, status
from sqlalchemy import func
from sqlalchemy.orm import Session

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer,
)

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

from app.models.contract import Contract
from app.models.obligation import Obligation
from app.models.renewal import Renewal
from app.models.user import User
from app.services.compliance_service import calculate_contract_compliance


CONTRACT_STATUSES = {
    "Draft",
    "Under Review",
    "Pending Approval",
    "Approved",
    "Active",
    "Expired",
    "Terminated",
}

OBLIGATION_STATUSES = {
    "Pending",
    "In Progress",
    "Completed",
    "Delayed",
    "Overdue",
}

RENEWAL_STATUSES = {
    "Upcoming",
    "In Progress",
    "Renewed",
    "Expired",
    "Cancelled",
}


def _validate_report_status(
    value,
    allowed_values,
    field_name
):
    if value is None:
        return

    if value not in allowed_values:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=(
                f"Invalid {field_name}. Allowed values: "
                + ", ".join(sorted(allowed_values))
            )
        )


def _validate_date_range(
    start_date,
    end_date
):
    if (
        start_date is not None
        and end_date is not None
        and start_date > end_date
    ):
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=(
                "start_date must be less than or equal to end_date"
            )
        )


# =========================================================
# REPORT ACCESS CONTROL
# =========================================================

def get_accessible_contracts(
    db: Session,
    current_user: User
):
    if current_user.role in {
        "Administrator",
        "Compliance Officer",
    }:
        return db.query(Contract)

    if current_user.role == "Legal Manager":
        return db.query(Contract).filter(
            Contract.created_by == current_user.id
        )

    if current_user.role in {
        "Contract Manager",
        "Employee",
    }:
        return db.query(Contract).filter(
            Contract.assigned_to == current_user.id
        )

    raise HTTPException(
        status_code=status.HTTP_403_FORBIDDEN,
        detail="You are not authorized to access reports"
    )


def get_accessible_contract_ids(
    db: Session,
    current_user: User
):
    query = get_accessible_contracts(
        db,
        current_user
    )

    return [
        contract_id
        for (contract_id,) in query.with_entities(
            Contract.id
        ).all()
    ]


# =========================================================
# CONTRACT ANALYTICS
# =========================================================

def get_contract_summary(
    db: Session,
    current_user: User
):
    query = get_accessible_contracts(
        db,
        current_user
    )

    total = query.with_entities(
        func.count(Contract.id)
    ).scalar() or 0

    active = query.filter(
        Contract.status == "Active"
    ).count()

    draft = query.filter(
        Contract.status == "Draft"
    ).count()

    under_review = query.filter(
        Contract.status.in_(
            [
                "Under Review",
                "Pending Approval",
            ]
        )
    ).count()

    approved = query.filter(
        Contract.status == "Approved"
    ).count()

    expired = query.filter(
        Contract.status == "Expired"
    ).count()

    terminated = query.filter(
        Contract.status == "Terminated"
    ).count()

    category_rows = (
        query.with_entities(
            Contract.category,
            func.count(Contract.id)
        )
        .group_by(Contract.category)
        .order_by(Contract.category)
        .all()
    )

    contracts_by_category = {
        category: count
        for category, count in category_rows
    }

    return {
        "total": total,
        "active": active,
        "draft": draft,
        "under_review": under_review,
        "approved": approved,
        "expired": expired,
        "terminated": terminated,
        "contracts_by_category": contracts_by_category,
    }


# =========================================================
# OBLIGATION ANALYTICS
# =========================================================

def get_obligation_summary(
    db: Session,
    current_user: User
):
    contract_ids = get_accessible_contract_ids(
        db,
        current_user
    )

    query = db.query(Obligation).filter(
        Obligation.contract_id.in_(contract_ids)
    )

    total = query.count()

    pending = query.filter(
        Obligation.status == "Pending"
    ).count()

    in_progress = query.filter(
        Obligation.status == "In Progress"
    ).count()

    completed = query.filter(
        Obligation.status == "Completed"
    ).count()

    delayed = query.filter(
        Obligation.status == "Delayed"
    ).count()

    overdue = query.filter(
        Obligation.status == "Overdue"
    ).count()

    return {
        "total": total,
        "pending": pending,
        "in_progress": in_progress,
        "completed": completed,
        "delayed": delayed,
        "overdue": overdue,
    }


# =========================================================
# RENEWAL ANALYTICS
# =========================================================

def get_renewal_summary(
    db: Session,
    current_user: User
):
    contract_ids = get_accessible_contract_ids(
        db,
        current_user
    )

    today = date.today()

    renewal_query = db.query(Renewal).filter(
        Renewal.contract_id.in_(contract_ids)
    )

    upcoming = renewal_query.filter(
        Renewal.renewal_date >= today,
        Renewal.status == "Upcoming"
    ).count()

    in_progress = renewal_query.filter(
        Renewal.status == "In Progress"
    ).count()

    renewed = renewal_query.filter(
        Renewal.status == "Renewed"
    ).count()

    expired = renewal_query.filter(
        Renewal.status == "Expired"
    ).count()

    cancelled = renewal_query.filter(
        Renewal.status == "Cancelled"
    ).count()

    expiry_limit = today + timedelta(days=30)

    approaching_expiry = (
        get_accessible_contracts(
            db,
            current_user
        )
        .filter(
            Contract.end_date.isnot(None),
            Contract.end_date >= today,
            Contract.end_date <= expiry_limit
        )
        .order_by(Contract.end_date)
        .all()
    )

    approaching_contracts = []

    for contract in approaching_expiry:

        days_remaining = (
            contract.end_date - today
        ).days

        approaching_contracts.append(
            {
                "contract_id": contract.id,
                "contract_number": contract.contract_number,
                "title": contract.title,
                "expiry_date": contract.end_date,
                "days_remaining": days_remaining,
            }
        )

    return {
        "upcoming": upcoming,
        "in_progress": in_progress,
        "renewed": renewed,
        "expired": expired,
        "cancelled": cancelled,
        "approaching_expiry": approaching_contracts,
    }


# =========================================================
# COMPLIANCE ANALYTICS
# =========================================================

def get_accessible_compliance(
    db: Session,
    current_user: User
):
    contracts = (
        get_accessible_contracts(
            db,
            current_user
        )
        .order_by(Contract.id)
        .all()
    )

    results = []

    for contract in contracts:

        result = calculate_contract_compliance(
            db,
            contract.id
        )

        result["contract_number"] = (
            contract.contract_number
        )

        results.append(result)

    return results


def get_compliance_report_summary(
    db: Session,
    current_user: User
):
    results = get_accessible_compliance(
        db,
        current_user
    )

    total_evaluated = len(results)

    compliant = sum(
        1
        for item in results
        if item["compliance_status"] == "Compliant"
    )

    pending = sum(
        1
        for item in results
        if item["compliance_status"] == "Pending"
    )

    delayed = sum(
        1
        for item in results
        if item["compliance_status"] == "Delayed"
    )

    non_compliant = sum(
        1
        for item in results
        if item["compliance_status"] == "Non-Compliant"
    )

    high_risk = sum(
        1
        for item in results
        if item["compliance_status"] == "High Risk"
    )

    scores = [
        item["compliance_score"]
        for item in results
        if item.get("compliance_score") is not None
    ]

    average_score = (
        round(
            sum(scores) / len(scores),
            2
        )
        if scores
        else 0.0
    )

    return {
        "total_evaluated": total_evaluated,
        "compliant": compliant,
        "pending": pending,
        "delayed": delayed,
        "non_compliant": non_compliant,
        "high_risk": high_risk,
        "average_score": average_score,
    }


# =========================================================
# RISK ANALYTICS
# =========================================================

def get_risk_summary(
    db: Session,
    current_user: User
):
    compliance_results = get_accessible_compliance(
        db,
        current_user
    )

    risk_contracts = []

    for item in compliance_results:

        if item["risk_level"] in (
            "Medium",
            "High",
        ):

            risk_contracts.append(
                {
                    "contract_id": item["contract_id"],
                    "contract_number": item["contract_number"],
                    "risk_level": item["risk_level"],
                    "overdue_obligations": item[
                        "overdue_obligations"
                    ],
                    "compliance_score": item[
                        "compliance_score"
                    ],
                }
            )

    return {
        "total_at_risk": len(risk_contracts),
        "contracts": risk_contracts,
    }


# =========================================================
# DASHBOARD SUMMARY
# =========================================================

def get_dashboard_summary(
    db: Session,
    current_user: User
):
    contract_summary = get_contract_summary(
        db,
        current_user
    )

    obligation_summary = get_obligation_summary(
        db,
        current_user
    )

    renewal_summary = get_renewal_summary(
        db,
        current_user
    )

    compliance_summary = get_compliance_report_summary(
        db,
        current_user
    )

    return {
        "contracts": {
            "total": contract_summary["total"],
            "active": contract_summary["active"],
            "draft": contract_summary["draft"],
            "under_review": contract_summary[
                "under_review"
            ],
            "approved": contract_summary[
                "approved"
            ],
            "expired": contract_summary["expired"],
            "terminated": contract_summary[
                "terminated"
            ],
        },

        "obligations": {
            "total": obligation_summary["total"],
            "pending": obligation_summary["pending"],
            "in_progress": obligation_summary[
                "in_progress"
            ],
            "completed": obligation_summary[
                "completed"
            ],
            "overdue": obligation_summary["overdue"],
        },

        "renewals": {
            "upcoming": renewal_summary["upcoming"],
            "in_progress": renewal_summary[
                "in_progress"
            ],
            "renewed": renewal_summary["renewed"],
            "expired": renewal_summary["expired"],
        },

        "compliance": {
            "compliant": compliance_summary[
                "compliant"
            ],
            "pending": compliance_summary["pending"],
            "delayed": compliance_summary["delayed"],
            "non_compliant": compliance_summary[
                "non_compliant"
            ],
            "high_risk": compliance_summary[
                "high_risk"
            ],
        },
    }


# =========================================================
# CONTRACT REPORT
# =========================================================

def generate_contract_report(
    db: Session,
    current_user: User,
    status_filter: str | None = None
):
    _validate_report_status(
        status_filter,
        CONTRACT_STATUSES,
        "contract status"
    )

    contracts_query = get_accessible_contracts(
        db,
        current_user
    )

    if status_filter is not None:

        contracts_query = contracts_query.filter(
            Contract.status == status_filter
        )

    contracts = (
        contracts_query
        .order_by(Contract.id)
        .all()
    )

    report_data = []

    for contract in contracts:

        assigned_user = None

        if contract.assigned_to:

            assigned_user = (
                db.query(User)
                .filter(
                    User.id == contract.assigned_to
                )
                .first()
            )

        report_data.append(
            {
                "contract_number":
                    contract.contract_number,

                "title":
                    contract.title,

                "category":
                    contract.category,

                "status":
                    contract.status,

                "start_date":
                    contract.start_date,

                "end_date":
                    contract.end_date,

                "assigned_user": (
                    assigned_user.full_name
                    if assigned_user
                    else None
                ),
            }
        )

    return {
        "report_type": "Contract Report",
        "total_records": len(report_data),
        "data": report_data,
    }


# =========================================================
# OBLIGATION REPORT
# =========================================================

def generate_obligation_report(
    db: Session,
    current_user: User,
    status_filter: str | None = None
):
    _validate_report_status(
        status_filter,
        OBLIGATION_STATUSES,
        "obligation status"
    )

    contract_ids = get_accessible_contract_ids(
        db,
        current_user
    )

    obligations_query = db.query(
        Obligation
    ).filter(
        Obligation.contract_id.in_(contract_ids)
    )

    if status_filter is not None:

        obligations_query = obligations_query.filter(
            Obligation.status == status_filter
        )

    obligations = (
        obligations_query
        .order_by(Obligation.id)
        .all()
    )

    report_data = []

    for obligation in obligations:

        contract = (
            db.query(Contract)
            .filter(
                Contract.id == obligation.contract_id
            )
            .first()
        )

        assigned_user = None

        if obligation.assigned_to:

            assigned_user = (
                db.query(User)
                .filter(
                    User.id == obligation.assigned_to
                )
                .first()
            )

        report_data.append(
            {
                "contract_number": (
                    contract.contract_number
                    if contract
                    else "N/A"
                ),

                "title":
                    obligation.title,

                "obligation_type":
                    obligation.obligation_type,

                "assigned_user": (
                    assigned_user.full_name
                    if assigned_user
                    else None
                ),

                "due_date":
                    obligation.due_date,

                "status":
                    obligation.status,

                "completion_date":
                    obligation.completion_date,
            }
        )

    return {
        "report_type": "Obligation Report",
        "total_records": len(report_data),
        "data": report_data,
    }


# =========================================================
# RENEWAL REPORT
# =========================================================

def generate_renewal_report(
    db: Session,
    current_user: User,
    status_filter: str | None = None,
    start_date: date | None = None,
    end_date: date | None = None
):
    _validate_report_status(
        status_filter,
        RENEWAL_STATUSES,
        "renewal status"
    )

    _validate_date_range(
        start_date,
        end_date
    )

    contract_ids = get_accessible_contract_ids(
        db,
        current_user
    )

    renewals_query = db.query(
        Renewal
    ).filter(
        Renewal.contract_id.in_(contract_ids)
    )

    if status_filter is not None:

        renewals_query = renewals_query.filter(
            Renewal.status == status_filter
        )

    if start_date is not None:

        renewals_query = renewals_query.filter(
            Renewal.renewal_date >= start_date
        )

    if end_date is not None:

        renewals_query = renewals_query.filter(
            Renewal.renewal_date <= end_date
        )

    renewals = (
        renewals_query
        .order_by(Renewal.id)
        .all()
    )

    report_data = []

    for renewal in renewals:

        contract = (
            db.query(Contract)
            .filter(
                Contract.id == renewal.contract_id
            )
            .first()
        )

        assigned_user = None

        if renewal.assigned_to:

            assigned_user = (
                db.query(User)
                .filter(
                    User.id == renewal.assigned_to
                )
                .first()
            )

        report_data.append(
            {
                "contract_number": (
                    contract.contract_number
                    if contract
                    else "N/A"
                ),

                "previous_expiry_date":
                    renewal.previous_expiry_date,

                "renewal_date":
                    renewal.renewal_date,

                "new_expiry_date":
                    renewal.new_expiry_date,

                "status":
                    renewal.status,

                "assigned_user": (
                    assigned_user.full_name
                    if assigned_user
                    else None
                ),
            }
        )

    return {
        "report_type": "Renewal Report",
        "total_records": len(report_data),
        "data": report_data,
    }


# =========================================================
# COMPLIANCE REPORT
# =========================================================

def generate_compliance_report(
    db: Session,
    current_user: User
):
    compliance_results = get_accessible_compliance(
        db,
        current_user
    )

    evaluation_date = date.today()

    report_data = []

    for item in compliance_results:

        report_data.append(
            {
                "contract_number":
                    item["contract_number"],

                "compliance_status":
                    item["compliance_status"],

                "compliance_score":
                    item["compliance_score"],

                "overdue_obligations":
                    item["overdue_obligations"],

                "risk_level":
                    item["risk_level"],

                "evaluation_date":
                    evaluation_date,
            }
        )

    return {
        "report_type": "Compliance Report",
        "total_records": len(report_data),
        "data": report_data,
    }


# =========================================================
# EXPORT HELPERS
# =========================================================

def _format_value(value):

    if value is None:
        return ""

    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")

    return str(value)


def _get_report_data(
    db: Session,
    current_user: User,
    report_type: str
):
    report_type = report_type.lower().strip()

    if report_type == "contract":

        return generate_contract_report(
            db,
            current_user
        )

    if report_type == "obligation":

        return generate_obligation_report(
            db,
            current_user
        )

    if report_type == "renewal":

        return generate_renewal_report(
            db,
            current_user
        )

    if report_type == "compliance":

        return generate_compliance_report(
            db,
            current_user
        )

    raise ValueError(
        "Invalid report type. "
        "Use contract, obligation, renewal, or compliance."
    )


def _get_report_columns(
    report_type: str
):
    report_type = report_type.lower().strip()

    if report_type == "contract":

        return [
            (
                "contract_number",
                "Contract Number"
            ),
            (
                "title",
                "Title"
            ),
            (
                "category",
                "Category"
            ),
            (
                "status",
                "Status"
            ),
            (
                "start_date",
                "Start Date"
            ),
            (
                "end_date",
                "End Date"
            ),
            (
                "assigned_user",
                "Assigned User"
            ),
        ]

    if report_type == "obligation":

        return [
            (
                "contract_number",
                "Contract Number"
            ),
            (
                "title",
                "Title"
            ),
            (
                "obligation_type",
                "Obligation Type"
            ),
            (
                "assigned_user",
                "Assigned User"
            ),
            (
                "due_date",
                "Due Date"
            ),
            (
                "status",
                "Status"
            ),
            (
                "completion_date",
                "Completion Date"
            ),
        ]

    if report_type == "renewal":

        return [
            (
                "contract_number",
                "Contract Number"
            ),
            (
                "previous_expiry_date",
                "Previous Expiry Date"
            ),
            (
                "renewal_date",
                "Renewal Date"
            ),
            (
                "new_expiry_date",
                "New Expiry Date"
            ),
            (
                "status",
                "Status"
            ),
            (
                "assigned_user",
                "Assigned User"
            ),
        ]

    if report_type == "compliance":

        return [
            (
                "contract_number",
                "Contract Number"
            ),
            (
                "compliance_status",
                "Compliance Status"
            ),
            (
                "compliance_score",
                "Compliance Score"
            ),
            (
                "overdue_obligations",
                "Overdue Obligations"
            ),
            (
                "risk_level",
                "Risk Level"
            ),
            (
                "evaluation_date",
                "Evaluation Date"
            ),
        ]

    raise ValueError(
        "Invalid report type. "
        "Use contract, obligation, renewal, or compliance."
    )


# =========================================================
# PDF EXPORT
# =========================================================

def generate_pdf_report(
    db: Session,
    current_user: User,
    report_type: str
):
    report = _get_report_data(
        db,
        current_user,
        report_type
    )

    columns = _get_report_columns(
        report_type
    )

    buffer = BytesIO()

    document = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
    )

    styles = getSampleStyleSheet()

    title_style = styles["Title"]

    normal_style = styles["Normal"]

    story = []

    story.append(
        Paragraph(
            report["report_type"],
            title_style
        )
    )

    story.append(
        Spacer(
            1,
            6 * mm
        )
    )

    story.append(
        Paragraph(
            f"Total Records: "
            f"{report['total_records']}",
            normal_style
        )
    )

    story.append(
        Paragraph(
            f"Generated Date: "
            f"{date.today().strftime('%Y-%m-%d')}",
            normal_style
        )
    )

    story.append(
        Spacer(
            1,
            5 * mm
        )
    )

    table_data = [
        [
            header
            for _, header in columns
        ]
    ]

    for item in report["data"]:

        row = []

        for key, _ in columns:

            value = _format_value(
                item.get(key)
            )

            row.append(
                Paragraph(
                    value,
                    styles["BodyText"]
                )
            )

        table_data.append(row)

    if len(table_data) == 1:

        table_data.append(
            [
                Paragraph(
                    "No records found.",
                    styles["BodyText"]
                )
                for _ in columns
            ]
        )

    table = Table(
        table_data,
        repeatRows=1
    )

    table.setStyle(
        TableStyle(
            [
                (
                    "BACKGROUND",
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#D9EAF7")
                ),
                (
                    "TEXTCOLOR",
                    (0, 0),
                    (-1, 0),
                    colors.black
                ),
                (
                    "FONTNAME",
                    (0, 0),
                    (-1, 0),
                    "Helvetica-Bold"
                ),
                (
                    "FONTSIZE",
                    (0, 0),
                    (-1, -1),
                    7
                ),
                (
                    "ALIGN",
                    (0, 0),
                    (-1, 0),
                    "CENTER"
                ),
                (
                    "VALIGN",
                    (0, 0),
                    (-1, -1),
                    "MIDDLE"
                ),
                (
                    "GRID",
                    (0, 0),
                    (-1, -1),
                    0.5,
                    colors.grey
                ),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [
                        colors.white,
                        colors.HexColor("#F7F7F7")
                    ]
                ),
                (
                    "LEFTPADDING",
                    (0, 0),
                    (-1, -1),
                    4
                ),
                (
                    "RIGHTPADDING",
                    (0, 0),
                    (-1, -1),
                    4
                ),
                (
                    "TOPPADDING",
                    (0, 0),
                    (-1, -1),
                    4
                ),
                (
                    "BOTTOMPADDING",
                    (0, 0),
                    (-1, -1),
                    4
                ),
            ]
        )
    )

    story.append(table)

    document.build(story)

    buffer.seek(0)

    return buffer


# =========================================================
# EXCEL EXPORT
# =========================================================

def generate_excel_report(
    db: Session,
    current_user: User,
    report_type: str
):
    report = _get_report_data(
        db,
        current_user,
        report_type
    )

    columns = _get_report_columns(
        report_type
    )

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Report"

    worksheet["A1"] = report["report_type"]

    worksheet["A1"].font = Font(
        bold=True,
        size=16
    )

    worksheet["A2"] = "Total Records"

    worksheet["B2"] = (
        report["total_records"]
    )

    worksheet["A3"] = "Generated Date"

    worksheet["B3"] = date.today()

    worksheet["B3"].number_format = (
        "yyyy-mm-dd"
    )

    header_row = 5

    for column_index, (_, header) in enumerate(
        columns,
        start=1
    ):

        cell = worksheet.cell(
            row=header_row,
            column=column_index,
            value=header
        )

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    for row_index, item in enumerate(
        report["data"],
        start=header_row + 1
    ):

        for column_index, (key, _) in enumerate(
            columns,
            start=1
        ):

            value = item.get(key)

            cell = worksheet.cell(
                row=row_index,
                column=column_index,
                value=value
            )

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

            if isinstance(value, date):

                cell.number_format = (
                    "yyyy-mm-dd"
                )

    worksheet.freeze_panes = "A6"

    for column_index, (_, header) in enumerate(
        columns,
        start=1
    ):

        max_length = len(header)

        for row_index in range(
            header_row + 1,
            header_row + 1 + len(
                report["data"]
            )
        ):

            cell_value = worksheet.cell(
                row=row_index,
                column=column_index
            ).value

            if cell_value is not None:

                value_length = len(
                    str(cell_value)
                )

                if value_length > max_length:

                    max_length = value_length

        adjusted_width = min(
            max_length + 2,
            40
        )

        worksheet.column_dimensions[
            get_column_letter(column_index)
        ].width = adjusted_width

    buffer = BytesIO()

    workbook.save(buffer)

    buffer.seek(0)

    return buffer