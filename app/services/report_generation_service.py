from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from sqlalchemy.orm import Session

from app.models.contract import Contract
from app.models.obligation import Obligation
from app.models.renewal import Renewal


REPORT_DIR = Path("generated_reports")
REPORT_DIR.mkdir(parents=True, exist_ok=True)


def _timestamp():
    return datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")


def generate_contract_report_excel(db: Session) -> str:
    filename = REPORT_DIR / f"contracts_{_timestamp()}.xlsx"

    contracts = db.query(Contract).order_by(Contract.id).all()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Contracts"

    headers = [
        "ID",
        "Contract Number",
        "Title",
        "Category",
        "Counterparty",
        "Start Date",
        "End Date",
        "Status",
        "Created By",
        "Assigned To",
    ]

    sheet.append(headers)

    for contract in contracts:
        sheet.append([
            contract.id,
            contract.contract_number,
            contract.title,
            contract.category,
            contract.counterparty_name,
            contract.start_date,
            contract.end_date,
            contract.status,
            contract.created_by,
            contract.assigned_to,
        ])

    workbook.save(filename)
    return str(filename)


def generate_obligation_report_excel(db: Session) -> str:
    filename = REPORT_DIR / f"obligations_{_timestamp()}.xlsx"

    obligations = db.query(Obligation).order_by(Obligation.id).all()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Obligations"

    headers = [
        "ID",
        "Contract ID",
        "Title",
        "Description",
        "Due Date",
        "Status",
        "Priority",
        "Responsible Party",
    ]

    sheet.append(headers)

    for obligation in obligations:
        sheet.append([
            obligation.id,
            obligation.contract_id,
            obligation.title,
            obligation.description,
            obligation.due_date,
            obligation.status,
            obligation.priority,
            obligation.responsible_party,
        ])

    workbook.save(filename)
    return str(filename)


def generate_renewal_report_excel(db: Session) -> str:
    filename = REPORT_DIR / f"renewals_{_timestamp()}.xlsx"

    renewals = db.query(Renewal).order_by(Renewal.id).all()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Renewals"

    headers = [
        "ID",
        "Contract ID",
        "Renewal Date",
        "Status",
        "Renewal Terms",
    ]

    sheet.append(headers)

    for renewal in renewals:
        sheet.append([
            renewal.id,
            renewal.contract_id,
            renewal.renewal_date,
            renewal.status,
            renewal.renewal_terms,
        ])

    workbook.save(filename)
    return str(filename)


def generate_contract_report_pdf(db: Session) -> str:
    filename = REPORT_DIR / f"contracts_{_timestamp()}.pdf"

    contracts = db.query(Contract).order_by(Contract.id).all()

    document = SimpleDocTemplate(
        str(filename),
        pagesize=A4,
    )

    styles = getSampleStyleSheet()
    elements = [
        Paragraph("ContractIQ - Contract Report", styles["Title"]),
        Spacer(1, 12),
        Paragraph(
            f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S')}",
            styles["Normal"],
        ),
        Spacer(1, 12),
    ]

    data = [[
        "ID",
        "Contract Number",
        "Title",
        "Counterparty",
        "Status",
        "End Date",
    ]]

    for contract in contracts:
        data.append([
            str(contract.id),
            str(contract.contract_number or ""),
            str(contract.title or ""),
            str(contract.counterparty_name or ""),
            str(contract.status or ""),
            str(contract.end_date or ""),
        ])

    table = Table(data, repeatRows=1)

    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))

    elements.append(table)
    document.build(elements)

    return str(filename)


def generate_dashboard_report_pdf(
    db: Session,
    dashboard_data: dict,
) -> str:
    filename = REPORT_DIR / f"dashboard_{_timestamp()}.pdf"

    document = SimpleDocTemplate(
        str(filename),
        pagesize=A4,
    )

    styles = getSampleStyleSheet()

    elements = [
        Paragraph("ContractIQ - Dashboard Report", styles["Title"]),
        Spacer(1, 12),
        Paragraph(
            f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S')}",
            styles["Normal"],
        ),
        Spacer(1, 20),
    ]

    rows = [["Metric", "Value"]]

    for section_name, section_data in dashboard_data.items():
        rows.append([section_name.title(), ""])

        for metric, value in section_data.items():
            rows.append([
                f"  {metric.replace('_', ' ').title()}",
                str(value),
            ])

    table = Table(rows, colWidths=[350, 120])

    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))

    elements.append(table)
    document.build(elements)

    return str(filename)
